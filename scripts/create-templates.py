"""Regenerate templates: uv run --with python-docx --with openpyxl --with python-pptx scripts/create-templates.py"""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]/'templates';ROOT.mkdir(exist_ok=True)
d=Document();d.core_properties.author='Cohnsoft';d.core_properties.title='Untitled';d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(11);d.add_paragraph();d.save(ROOT/'blank.docx')
w=Workbook();w.active.title='Sheet 1';w.properties.creator='Cohnsoft';w.save(ROOT/'blank.xlsx')
p=Presentation();p.core_properties.author='Cohnsoft';p.slides.add_slide(p.slide_layouts[6]);p.save(ROOT/'blank.pptx')
print('Created blank DOCX, XLSX and PPTX templates.')
